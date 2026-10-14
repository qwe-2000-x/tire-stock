"""
把两份价格表合成一份 seed.sql，用于初始化数据库。

  npm run db:seed          灌到本地
  npm run db:seed:remote   灌到线上

价格表名称取自文件名（和网页里「导入价格表」自动填的名字一致），
这样以后在网页上重新导入同名文件时，会正确替换掉这里灌进去的数据。

以后换了新的价格表文件，只要改下面 SOURCES 里的文件名，重跑这个脚本就行。
"""

import sys
import pathlib

sys.path.insert(0, str(pathlib.Path(__file__).parent / "pylibs"))
sys.stdout.reconfigure(encoding="utf-8")

import xlrd
import openpyxl

ROOT = pathlib.Path(__file__).resolve().parent.parent
WORK = ROOT.parent

# (文件, 价格表名称, 格式, 表头在第几行(从1算), 各列的 0 基下标)
#
# 价格表名称就是网页里「导入价格表」要填的名字。
# 以后换新表格时，把文件名改成一样的（安睿驰价格表.xls / 瑞途价格表.xlsx），
# 在网页上导入时会自动替换掉对应的那一份。
SOURCES = [
    {
        "file": WORK / "安睿驰价格表.xls",
        "list": "安睿驰价格表",
        "format": "xls",
        "header_row": 2,
        "cols": {"model": 1, "series": 2, "pattern": 3, "store": 4, "net": 5},
    },
    {
        "file": WORK / "瑞途价格表.xlsx",
        "list": "瑞途价格表",
        "format": "xlsx",
        "header_row": 3,
        "cols": {"model": 1, "series": None, "pattern": 2, "store": 3, "net": 4},
    },
]


def txt(v):
    if v is None:
        return ""
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    return str(v).strip()


def num(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace(",", "").replace("¥", "").replace("元", "").strip()
    try:
        return float(s)
    except ValueError:
        return None


def read_rows(src):
    path = src["file"]
    if src["format"] == "xls":
        sheet = xlrd.open_workbook(str(path)).sheet_by_index(0)
        return [[sheet.cell_value(r, c) for c in range(sheet.ncols)] for r in range(sheet.nrows)]
    ws = openpyxl.load_workbook(str(path), data_only=True).worksheets[0]
    return [
        [ws.cell(row=r + 1, column=c + 1).value for c in range(ws.max_column)]
        for r in range(ws.max_row)
    ]


def sql_str(value):
    if value is None or value == "":
        return "NULL"
    return "'" + str(value).replace("'", "''") + "'"


def sql_num(value):
    return "NULL" if value is None else repr(float(value))


def main():
    lines = [
        "-- 由 tools/build_seed.py 自动生成，请勿手改",
        "-- 内容：两份价格表的全部商品，导入前会先清空 products 表",
        "",
        "DELETE FROM products;",
        "",
    ]
    total = 0
    for src in SOURCES:
        rows = read_rows(src)
        cols = src["cols"]
        count = 0
        for row in rows[src["header_row"]:]:
            model = txt(row[cols["model"]] if cols["model"] < len(row) else "")
            if not model:
                continue
            pattern = txt(row[cols["pattern"]]) if cols["pattern"] is not None and cols["pattern"] < len(row) else ""
            series = txt(row[cols["series"]]) if cols["series"] is not None and cols["series"] < len(row) else ""
            store = num(row[cols["store"]]) if cols["store"] is not None and cols["store"] < len(row) else None
            net = num(row[cols["net"]]) if cols["net"] is not None and cols["net"] < len(row) else None
            lines.append(
                "INSERT INTO products (model, pattern, series, storePrice, netPrice, priceList, updatedAt) "
                f"VALUES ({sql_str(model)}, {sql_str(pattern)}, {sql_str(series)}, "
                f"{sql_num(store)}, {sql_num(net)}, {sql_str(src['list'])}, datetime('now'));"
            )
            count += 1
        total += count
        print(f"{src['file'].name} → 「{src['list']}」{count} 行")

    lines.append("")
    out = ROOT / "seed.sql"
    out.write_text("\n".join(lines), encoding="utf-8")
    print(f"合计 {total} 行，已写入 {out.name}")


if __name__ == "__main__":
    main()
